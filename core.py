"""Bounded tools, local document retrieval, and conservative escalation checks."""
import io
import math
import re
from pathlib import Path

EMERGENCY_MESSAGE = ('Some of your words may describe an urgent situation. If these symptoms are happening now, '
    'contact your local emergency service or go to the nearest emergency department now. Do not wait for this chatbot. '
    'Ask someone nearby to help and follow emergency-dispatch instructions. '
    'This keyword check cannot establish whether you are safe or diagnose a condition.')
PATTERNS = [r'chest pain', r'chest pressure', r'(cannot|can.t|unable to) breathe', r'severe (breath|bleed)',
    r'(face|facial) droop', r'sudden.*(weakness|speech)', r'unconscious', r'overdose', r'not breathing',
    r'kill myself', r'suicid', r'self.harm', r'سینے میں درد', r'سانس نہیں', r'بے ہوش']


def urgent_flag(text):
    """Conservative, non-diagnostic alert. Negated/historical mentions may also trigger."""
    return any(re.search(p, text, re.I) for p in PATTERNS)


def bmi(weight_kg, height_cm):
    weight_kg, height_cm = float(weight_kg), float(height_cm)
    if not (math.isfinite(weight_kg) and math.isfinite(height_cm) and 1 <= weight_kg <= 500 and 30 <= height_cm <= 260):
        raise ValueError('Enter plausible weight (1–500 kg) and height (30–260 cm).')
    return {'bmi': round(weight_kg / (height_cm / 100) ** 2, 1),
            'note': 'Arithmetic only; no age-specific, pregnancy or clinical interpretation is applied.'}


def chunks(text, source, page=None):
    words = text.split()
    return [{'source': source, 'page': page, 'text': ' '.join(words[i:i+220])}
            for i in range(0, len(words), 180) if words[i:i+220]]


def extract(name, data):
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Maximum document size is 10 MB.')
    ext = Path(name).suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted:
            raise ValueError('Unlock the PDF locally first.')
        if len(pdf.pages) > 150:
            raise ValueError('Maximum 150 PDF pages per document.')
        result = [c for i, page in enumerate(pdf.pages, 1) for c in chunks(page.extract_text() or '', name, i)]
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
        result = chunks(text, name)
    elif ext in ('.txt', '.md'):
        text = data.decode('utf-8-sig')
        if text.startswith('%PDF-') or '\x00' in text or text.startswith('PK\x03\x04'):
            raise ValueError('Binary file detected; use the correct extension.')
        result = chunks(text, name)
    else:
        raise ValueError('Supported: PDF, DOCX, TXT, Markdown. Images/scans require OCR first.')
    if not result or sum(len(c['text']) for c in result) < 20:
        raise ValueError('No readable text found. Scanned/image documents require OCR first.')
    return result


def retrieve(query, passages, limit=4):
    tokens = set(re.findall(r'\w{3,}', query.casefold())) - {'the', 'and', 'what', 'how', 'with', 'for', 'does', 'are'}
    ranked = []
    for i, passage in enumerate(passages):
        words = re.findall(r'\w{3,}', passage['text'].casefold())
        score = sum(math.log1p(words.count(t)) for t in tokens)
        if score:
            ranked.append((score, i, passage))
    ranked.sort(key=lambda item: (-item[0], item[1]))
    return [dict(p, citation=f'S{i+1}') for _, i, p in ranked[:limit]]


def handoff(concern='', duration='', medicines='', allergies='', questions=''):
    return {'type': 'UNREVIEWED handoff draft — not sent', 'main_concern': str(concern)[:3000],
            'duration': str(duration)[:500], 'reported_medicines': str(medicines)[:2000],
            'reported_allergies': str(allergies)[:1000], 'questions_for_clinician': str(questions)[:2000],
            'next_step': 'Check facts and share with a licensed clinician yourself. No appointment or message has been sent.'}
